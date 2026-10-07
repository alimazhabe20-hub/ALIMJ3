"""v77_platform: knowledge responsibilities."""
from .v77_platform_common import *  # noqa: F401,F403
from . import v77_platform_common as _common
globals().update({k:v for k,v in _common.__dict__.items() if not k.startswith('__')})


def extract_document(path: str|Path) -> dict[str,Any]:
    p=Path(path)
    if not p.exists() or not p.is_file():return {"ok":False,"error":"missing"}
    ext=p.suffix.lower();text="";tables=[]
    try:
        if ext==".pdf":
            from pypdf import PdfReader
            reader=PdfReader(str(p));text="\n".join((page.extract_text() or "") for page in reader.pages)
        elif ext==".docx":
            from docx import Document
            doc=Document(str(p));text="\n".join(x.text for x in doc.paragraphs if x.text.strip())
            tables=[[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
        elif ext in {".txt",".md",".csv"}:
            text=p.read_text(encoding="utf-8",errors="replace")
            if ext==".csv":
                rows=list(csv.reader(io.StringIO(text)));tables=rows[:200]
        else:return {"ok":False,"error":"unsupported_type"}
        return {"ok":True,"type":ext,"text":text[:200000],"characters":len(text),"tables":tables[:50],"sha256":backup_fingerprint(p)["sha256"]}
    except Exception:return {"ok":False,"error":"document_parse_failed"}

def register_plugin(name: str, version: str, handler: Callable[...,Any], permissions: Iterable[str]=(), trusted: bool=False) -> dict[str,Any]:
    if not name or not callable(handler):return {"ok":False,"error":"invalid_plugin"}
    if not trusted:return {"ok":False,"error":"trust_required"}
    perms={str(x) for x in permissions if str(x) in {"read","network","files","market","admin"}}
    _PLUGINS[str(name)]={"version":str(version),"handler":handler,"permissions":perms,"trusted":True}
    return {"ok":True,"name":str(name),"version":str(version),"permissions":sorted(perms)}

def plugin_snapshot() -> dict[str,Any]:
    return {k:{"version":v["version"],"permissions":sorted(v["permissions"]),"trusted":v["trusted"]} for k,v in _PLUGINS.items()}

def update_conversation(user_id: int, role: str, content: str, metadata: dict[str,Any]|None=None) -> None:
    if role not in {"user","assistant","tool","system"}:role="user"
    c=_db();c.execute("INSERT INTO v77_conversation(id,user_id,role,content,metadata_json) VALUES(?,?,?,?,?)",(uuid.uuid4().hex,int(user_id),role,redact(content,10000),json.dumps(metadata or {},ensure_ascii=False)[:6000]));c.execute("DELETE FROM v77_conversation WHERE user_id=? AND id NOT IN (SELECT id FROM v77_conversation WHERE user_id=? ORDER BY created_at DESC LIMIT ?)",(int(user_id),int(user_id),MAX_STATE_TURNS));c.commit();c.close()

def get_conversation(user_id: int, limit: int=MAX_STATE_TURNS) -> list[dict[str,Any]]:
    c=_db();rows=c.execute("SELECT role,content,metadata_json,created_at FROM v77_conversation WHERE user_id=? ORDER BY created_at DESC LIMIT ?",(int(user_id),max(1,min(MAX_STATE_TURNS,int(limit))))).fetchall();c.close();out=[]
    for r in reversed(rows):
        try:m=json.loads(r[2] or "{}")
        except Exception:m={}
        out.append({"role":r[0],"content":r[1],"metadata":m,"created_at":r[3]})
    return out

def generate_report(title: str, rows: list[dict[str,Any]], fmt: str="json") -> tuple[bytes,str,str]:
    safe=re.sub(r"[^\w\- ]+","_",title or "report")[:80] or "report";fmt=fmt.lower();keys=sorted({k for r in rows for k in r})
    if fmt=="json":return json.dumps({"title":title,"rows":rows},ensure_ascii=False,indent=2).encode(),safe+".json","application/json"
    if fmt=="csv":
        s=io.StringIO();w=csv.DictWriter(s,fieldnames=keys);w.writeheader();w.writerows(rows);return s.getvalue().encode("utf-8-sig"),safe+".csv","text/csv"
    if fmt=="xlsx":
        try:
            from openpyxl import Workbook
            wb=Workbook();ws=wb.active;ws.title=(title or "Report")[:31];ws.append(keys)
            for r in rows:ws.append([r.get(k,"") for k in keys])
            b=io.BytesIO();wb.save(b);return b.getvalue(),safe+".xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        except Exception:return generate_report(title,rows,"csv")
    if fmt=="docx":
        try:
            from docx import Document
            doc=Document();doc.add_heading(title or "Report",0);table=doc.add_table(rows=1,cols=max(1,len(keys))); 
            for i,k in enumerate(keys):table.rows[0].cells[i].text=str(k)
            for r in rows:
                cells=table.add_row().cells
                for i,k in enumerate(keys):cells[i].text=str(r.get(k,""))
            b=io.BytesIO();doc.save(b);return b.getvalue(),safe+".docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except Exception:return generate_report(title,rows,"json")
    if fmt=="pdf":
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle
            from reportlab.lib import colors
            from reportlab.lib.styles import getSampleStyleSheet
            b=io.BytesIO();doc=SimpleDocTemplate(b,pagesize=A4);styles=getSampleStyleSheet();story=[Paragraph(title or "Report",styles["Title"])];data=[keys]+[[str(r.get(k,"")) for k in keys] for r in rows]
            table=Table(data,repeatRows=1);table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.25,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.lightgrey)]));story.append(table);doc.build(story);return b.getvalue(),safe+".pdf","application/pdf"
        except Exception:return generate_report(title,rows,"json")
    return generate_report(title,rows,"json")

def rank_sources(items: Iterable[dict[str,Any]], query: str="") -> list[dict[str,Any]]:
    seen=set();out=[];terms=set(re.findall(r"\w{3,}",query.lower()))
    for raw in list(items)[:200]:
        x=dict(raw);url=str(x.get("url","")).strip();title=str(x.get("title","")).strip();key=hashlib.sha256((url or title).lower().encode()).hexdigest()
        if key in seen:continue
        seen.add(key);text=(title+" "+str(x.get("content", ""))).lower();match=sum(t in text for t in terms);trust=float(x.get("trust",0) or 0)+(0.1 if url.startswith("https://") else 0);x.update(score=round(match+trust,3),untrusted=True);out.append(x)
    return sorted(out,key=lambda x:x["score"],reverse=True)

def verify_claim(claim: str, evidence: Iterable[str]) -> dict[str,Any]:
    tokens=set(re.findall(r"\w{4,}",str(claim).lower()));ev=" ".join(map(str,evidence)).lower();matched=sum(t in ev for t in tokens);confidence=matched/max(1,len(tokens));return {"claim":redact(claim,1500),"supported":confidence>=.5,"confidence":round(confidence,3),"matched_terms":matched,"terms":len(tokens)}

def research_pack(query: str, sources: Iterable[dict[str,Any]]) -> dict[str,Any]:
    ranked=rank_sources(sources,query);claims=[]
    for item in ranked[:8]:
        content=str(item.get("content","") or item.get("snippet","")).strip()
        if content:claims.append(verify_claim(query,[content]))
    return {"query":redact(query,1000),"sources":ranked[:10],"evidence_checks":claims,"confidence":round(sum(x["confidence"] for x in claims)/len(claims),3) if claims else 0}
