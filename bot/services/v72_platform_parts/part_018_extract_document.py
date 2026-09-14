from typing import Any

# Auto-split part 18: extract_document
def extract_document(data: bytes, filename: str, mime: str = "") -> dict[str, Any]:
    if len(data) > MAX_DOCUMENT_BYTES: raise ValueError("too_large")
    name = Path(filename or "file").name[:200]
    ext = Path(name).suffix.lower()
    kind = "text"
    text = ""
    meta: dict[str, Any] = {"name": name, "size": len(data), "mime": mime}
    if ext in {".txt", ".md", ".csv", ".json", ".py", ".js", ".html", ".xml", ".log"} or mime.startswith("text/"):
        text = data.decode("utf-8", errors="replace")[:MAX_DOCUMENT_CHARS]
        kind = "csv" if ext == ".csv" else "text"
        if ext == ".csv":
            rows = list(csv.reader(io.StringIO(text)))[:100]
            meta["rows"] = len(rows); meta["columns"] = max((len(r) for r in rows), default=0)
    elif ext == ".json" or mime == "application/json":
        raw = data.decode("utf-8", errors="replace")
        try:
            obj = json.loads(raw); text = json.dumps(obj, ensure_ascii=False, indent=2)[:MAX_DOCUMENT_CHARS]; kind = "json"
        except Exception: text = raw[:MAX_DOCUMENT_CHARS]
    elif ext == ".pdf" or mime == "application/pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        meta["pages"] = len(reader.pages)
        chunks = []
        for i, page in enumerate(reader.pages[:50]):
            t = page.extract_text() or ""
            if t.strip(): chunks.append(f"--- صفحه {i+1} ---\n{t}")
        text = "\n\n".join(chunks)[:MAX_DOCUMENT_CHARS]; kind = "pdf"
    elif ext == ".docx" or "wordprocessingml" in mime:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if any(not _safe_member(n) for n in z.namelist()): raise ValueError("unsafe_archive")
            xml = z.read("word/document.xml").decode("utf-8", errors="ignore")
        text = re.sub(r"<[^>]+>", " ", xml)
        text = re.sub(r"\s+", " ", text)[:MAX_DOCUMENT_CHARS]; kind = "docx"
    elif ext in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheets = []
            for ws in wb.worksheets[:10]:
                rows = list(ws.iter_rows(values_only=True))[:100]
                sheets.append(f"--- {ws.title} ---\n" + "\n".join(", ".join("" if v is None else str(v) for v in row) for row in rows))
            text = "\n\n".join(sheets)[:MAX_DOCUMENT_CHARS]; meta["sheets"] = len(wb.worksheets); kind = "spreadsheet"
        except ImportError as exc: raise ValueError("spreadsheet_dependency_missing") from exc
    else:
        raise ValueError("unsupported_document")
    digest = hashlib.sha256(data).hexdigest()
    return {"name": name, "size": len(data), "sha256": digest, "kind": kind, "text": text, "meta": meta}
