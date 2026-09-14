from pathlib import Path
from typing import Any

# Auto-split part 47: extract_document
def extract_document(path: str|Path) -> dict[str,Any]:
    p=Path(path)
    if not p.exists() or not p.is_file():return {"ok":False,"error":"missing"}
    ext=p.suffix.lower();text="";tables=[]
    try:
        if ext==".pdf":
            from pypdf import PdfReader
            reader=PdfReader(str(p));text="\n".join((page.extract_text() or "") for page in reader.pages)
        elif ext==".docx":
            from docx import Document
            doc=Document(str(p));text="\n".join(x.text for x in doc.paragraphs if x.text.strip())
            tables=[[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
        elif ext in {".txt",".md",".csv"}:
            text=p.read_text(encoding="utf-8",errors="replace")
            if ext==".csv":
                rows=list(csv.reader(io.StringIO(text)));tables=rows[:200]
        else:return {"ok":False,"error":"unsupported_type"}
        return {"ok":True,"type":ext,"text":text[:200000],"characters":len(text),"tables":tables[:50],"sha256":backup_fingerprint(p)["sha256"]}
    except Exception:return {"ok":False,"error":"document_parse_failed"}
