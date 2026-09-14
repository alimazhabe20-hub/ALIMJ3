from typing import Any

# Auto-split part 52: generate_report
def generate_report(title: str, rows: list[dict[str,Any]], fmt: str="json") -> tuple[bytes,str,str]:
    safe=re.sub(r"[^\w\- ]+","_",title or "report")[:80] or "report";fmt=fmt.lower();keys=sorted({k for r in rows for k in r})
    if fmt=="json":return json.dumps({"title":title,"rows":rows},ensure_ascii=False,indent=2).encode(),safe+".json","application/json"
    if fmt=="csv":
        s=io.StringIO();w=csv.DictWriter(s,fieldnames=keys);w.writeheader();w.writerows(rows);return s.getvalue().encode("utf-8-sig"),safe+".csv","text/csv"
    if fmt=="xlsx":
        try:
            from openpyxl import Workbook
            wb=Workbook();ws=wb.active;ws.title=(title or "Report")[:31];ws.append(keys)
            for r in rows:ws.append([r.get(k,"") for k in keys])
            b=io.BytesIO();wb.save(b);return b.getvalue(),safe+".xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        except Exception:return generate_report(title,rows,"csv")
    if fmt=="docx":
        try:
            from docx import Document
            doc=Document();doc.add_heading(title or "Report",0);table=doc.add_table(rows=1,cols=max(1,len(keys))); 
            for i,k in enumerate(keys):table.rows[0].cells[i].text=str(k)
            for r in rows:
                cells=table.add_row().cells
                for i,k in enumerate(keys):cells[i].text=str(r.get(k,""))
            b=io.BytesIO();doc.save(b);return b.getvalue(),safe+".docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except Exception:return generate_report(title,rows,"json")
    if fmt=="pdf":
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle
            from reportlab.lib import colors
            from reportlab.lib.styles import getSampleStyleSheet
            b=io.BytesIO();doc=SimpleDocTemplate(b,pagesize=A4);styles=getSampleStyleSheet();story=[Paragraph(title or "Report",styles["Title"])];data=[keys]+[[str(r.get(k,"")) for k in keys] for r in rows]
            table=Table(data,repeatRows=1);table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.25,colors.grey),("BACKGROUND",(0,0),(-1,0),colors.lightgrey)]));story.append(table);doc.build(story);return b.getvalue(),safe+".pdf","application/pdf"
        except Exception:return generate_report(title,rows,"json")
    return generate_report(title,rows,"json")
