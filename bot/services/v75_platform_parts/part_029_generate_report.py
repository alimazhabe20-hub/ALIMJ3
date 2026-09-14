from typing import Any

# Auto-split part 29: generate_report
def generate_report(title:str, rows:list[dict[str,Any]], fmt:str="json") -> tuple[bytes,str,str]:
    fmt=fmt.lower().strip()
    safe_title=re.sub(r"[^\w\- ]+","_",title or "report")[:80] or "report"
    if fmt=="json": return json.dumps({"title":title,"rows":rows},ensure_ascii=False,indent=2).encode(),safe_title+".json","application/json"
    if fmt=="csv":
        keys=sorted({k for r in rows for k in r})
        buf=io.StringIO(); w=csv.DictWriter(buf,fieldnames=keys); w.writeheader(); w.writerows(rows)
        return buf.getvalue().encode("utf-8-sig"),safe_title+".csv","text/csv"
    if fmt=="xlsx":
        try:
            from openpyxl import Workbook
            wb=Workbook(); ws=wb.active; ws.title="Report"; keys=sorted({k for r in rows for k in r}); ws.append(keys)
            for r in rows: ws.append([r.get(k,"") for k in keys])
            out=io.BytesIO(); wb.save(out); return out.getvalue(),safe_title+".xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        except Exception: return generate_report(title,rows,"csv")
    if fmt=="docx":
        try:
            from docx import Document
            d=Document(); d.add_heading(title or "Report",0)
            for r in rows: d.add_paragraph(" | ".join(f"{k}: {v}" for k,v in r.items()))
            out=io.BytesIO(); d.save(out); return out.getvalue(),safe_title+".docx","application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except Exception: return generate_report(title,rows,"json")
    if fmt=="pdf":
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            out=io.BytesIO(); doc=SimpleDocTemplate(out,pagesize=A4); styles=getSampleStyleSheet(); story=[Paragraph(title or "Report",styles["Title"])]
            for r in rows: story.extend([Paragraph(redact(" | ".join(f"{k}: {v}" for k,v in r.items()),1800),styles["BodyText"]),Spacer(1,8)])
            doc.build(story); return out.getvalue(),safe_title+".pdf","application/pdf"
        except Exception: return generate_report(title,rows,"json")
    return generate_report(title,rows,"json")
